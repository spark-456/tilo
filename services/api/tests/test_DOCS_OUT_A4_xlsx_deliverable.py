"""DOCS-OUT A4: an Excel deliverable flows from the approved brief to a downloadable .xlsx."""

import asyncio
import io
import json

import pytest
from openpyxl import load_workbook
from test_DOCS_OUT_A3_office_deliverables import WORLD, XLSX, _case  # noqa: F401
from tilo_api import report_render
from tilo_harness.deliverable import DeliverableRefused, render_format
from tilo_harness.outputs import RenderedOffice


def _spec(rows: int = 3, tables: int = 1) -> str:
    return json.dumps(
        {
            "title": "Budget",
            "date": "2026-10-16",
            "tables": [
                {
                    "id": f"t{n}",
                    "columns": ["Item", "Due", "Cost"],
                    "rows": [[f"item {i}", "2026-11-01", i * 1.5] for i in range(rows)],
                }
                for n in range(tables)
            ],
            "sections": [{"heading": "Costs", "blocks": [{"type": "table", "table_id": "t0"}]}],
        }
    )


def test_DOCS_OUT_A4_an_xlsx_brief_selects_the_workbook_render():
    unit = {"role": "writer", "unit_type": "deliver"}
    assert render_format(unit, {"deliverables": [{"format": "xlsx"}]}) == "xlsx"
    assert render_format({"role": "specialist", "unit_type": "deliver"}, {"deliverables": [{"format": "xlsx"}]}) is None


def test_DOCS_OUT_A4_report_render_makes_a_workbook_with_more_rows_than_the_pdf_door_takes():
    output = asyncio.run(report_render.render_report(_spec(rows=400), fmt="xlsx"))  # 1200 cells would stop a door
    assert type(output) is RenderedOffice and output.media_type == XLSX
    sheet = load_workbook(io.BytesIO(output.content))["Costs"]
    assert sheet.max_row == 402 and sheet["C401"].value == 399 * 1.5 and sheet["C402"].value == "=SUM(C2:C401)"
    assert json.loads(output.spec_text)["title"] == "Budget"


def test_DOCS_OUT_A4_a_workbook_over_the_cell_cap_is_refused_with_a_fixed_reason():
    with pytest.raises(DeliverableRefused) as caught:
        asyncio.run(report_render.render_report(_spec(rows=500, tables=8), fmt="xlsx"))
    assert caught.value.code in {"too_large", "render_failed"}


def test_DOCS_OUT_A4_the_result_route_downloads_the_workbook_with_an_xlsx_name():
    body = asyncio.run(report_render.render_report(_spec(), fmt="xlsx")).content
    response = _case(XLSX, body).get(f"/worlds/{WORLD}/result")
    assert response.status_code == 200 and response.content == body
    assert response.headers["content-type"].split(";")[0] == XLSX
    assert response.headers["content-disposition"] == 'attachment; filename="tilo-result.xlsx"'
