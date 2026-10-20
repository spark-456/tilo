"""DEMO-DELIVER: every deliverable type Tilo promises comes out as a real file on the live path, end to end.

One writer reply per format (a fake provider; no paid call) goes through the real `ToolUnitRunner` with the API's
`render_report` door (as `run_tools` wires it), `outputs.keep` stores it in the run's artifact store (the Office screen
included), and the owner's `GET /worlds/{id}/result` serves it. Each body is then parsed back: the PDF by pypdf, the
Word file's XML, the deck by python-pptx, the workbook by openpyxl, the CSV by `csv`, and the website by the WS checker.
"""

import asyncio
import csv
import io
import json
import zipfile
from datetime import UTC, date, datetime
from functools import partial
from xml.etree import ElementTree as ET

import pytest
from _prod_wire_2b_lib import FakeBucket
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
from tilo_api import deps, report_render
from tilo_api.office_screen import office_text
from tilo_api.results import result_store, router
from tilo_api.run_stores import GcsBlobBackend
from tilo_harness.executor import UnitContext
from tilo_harness.outputs import keep
from tilo_harness.tool_runner import ToolUnitRunner
from tilo_memory.artifacts import ArtifactStore, MemoryMetaBackend, Scope
from tilo_memory.website_check import check_website
from tilo_providers.fake import FakeModelProvider
from tilo_providers.interfaces import GenerateRequest, GenerateResponse
from tilo_providers.safety import FakeSafetyProvider

OWNER, WORLD, RUN = "u_owner", "w_deliver", "r_deliver"
UNIT = {"id": "U5", "title": "World time zones", "unit_type": "deliver", "role": "writer"}
TODAY = date(2026, 10, 20)
ROWS = [
    ["UTC-08:00 Pacific", -8, "Los Angeles, Vancouver", 2, "West coast"],
    ["UTC-05:00 Eastern", -5, "New York, Toronto", 2, "=1+1 stays text"],
    ["UTC+00:00 Greenwich", 0, "London, Lisbon, Accra", 3, "Winter time"],
    ["UTC+05:30 India", 5.5, "Mumbai, Delhi, Kolkata", 3, "Half-hour offset"],
    ["UTC+09:00 Japan", 9, "Tokyo, Osaka", 2, "No daylight saving"],
]
SPEC = {
    "title": "World time zones and cities",
    "subtitle": "Standard offsets from UTC",
    "date": "1999-01-01",  # the code's date wins
    "tables": [
        {
            "id": "zones",
            "columns": ["Time zone", "UTC offset (hours)", "Major cities", "Cities listed", "Note"],
            "rows": ROWS,
            "formulas": [{"column": "Offset (minutes)", "expr": "[UTC offset (hours)] * 60"}],
        }
    ],
    "charts": [
        {"id": "c1", "kind": "bar", "table_id": "zones", "x": "Time zone", "y": "Cities listed", "title": "Cities"}
    ],
    "sections": [
        {
            "heading": "Overview",
            "blocks": [
                {"type": "paragraph", "text": "Five zones that cover the largest cities named in the research."},
                {"type": "bullets", "bullets": ["India keeps a half-hour offset.", "Japan has no daylight saving."]},
            ],
        },
        {"heading": "Zones", "blocks": [{"type": "table", "table_id": "zones"}, {"type": "chart", "chart_id": "c1"}]},
    ],
}
PAGE = (
    '<!doctype html>\n<html lang=en><head><meta charset="utf-8">'
    '<meta name="viewport" content="width=device-width, initial-scale=1"><title>Time zones</title>'
    "<style>body{font-family:system-ui,sans-serif}</style></head><body>\n<h1>World clocks</h1>\n"
    '<label>Zone <select id="z"><option>Asia/Tokyo</option><option>Europe/London</option></select></label>\n'
    '<button type="button" id="go">Show</button><p id="out"></p>\n'
    "<script>\ndocument.getElementById('go').addEventListener('click', function () {\n"
    "  var zone = document.getElementById('z').value;\n"
    "  document.getElementById('out').textContent = new Date().toLocaleTimeString('en-GB', {timeZone: zone});\n"
    '  window.scrollTo({ top: 0, behavior: "smooth" });\n'
    "});\n</script>\n</body></html>"
)
OFFICE = "application/vnd.openxmlformats-officedocument."
MIME = {
    "pdf": "application/pdf",
    "docx": OFFICE + "wordprocessingml.document",
    "pptx": OFFICE + "presentationml.presentation",
    "xlsx": OFFICE + "spreadsheetml.sheet",
    "csv": "text/csv",
}


class Writer(FakeModelProvider):
    def __init__(self, *replies: str) -> None:
        super().__init__("script")
        self.replies, self.requests = list(replies), []

    def generate(self, request: GenerateRequest) -> GenerateResponse:
        self.requests.append(request)
        return GenerateResponse(text=self.replies.pop(0), profile_id="script")


def artifact_store() -> ArtifactStore:
    return ArtifactStore(
        MemoryMetaBackend(),
        GcsBlobBackend(FakeBucket()),
        clock=lambda: datetime(2026, 10, 20, tzinfo=UTC),
        redactor=lambda text: text,
        archive_lookup=lambda *_: None,
        alert_sink=lambda _: None,
        screener=FakeSafetyProvider(),
    ).with_upload_text(office_text)  # the run's store screens a rendered Office file as its text


def deliver(fmt: str, *replies: str) -> tuple[TestClient, Writer, object]:
    """Runs the deliver unit for `fmt`, keeps its one output and returns a client for the owner's result route."""
    writer = Writer(*replies)
    runner = ToolUnitRunner(
        {"writer": writer},
        today=lambda: TODAY,
        report_renderer=partial(report_render.render_report, fmt=fmt, template=None, report_date=TODAY),
    )
    brief = {"goal": "All time zones and the cities in them", "deliverables": [{"title": "Time zones", "format": fmt}]}
    outcome = asyncio.run(runner.run(UNIT, UnitContext("A_w2", (), brief=brief)))
    assert outcome.success and len(outcome.outputs) == 1, outcome.summary
    artifacts, scope = artifact_store(), Scope(deps.ORG_ID, WORLD, OWNER)
    kept = keep(
        list(outcome.outputs), artifacts=artifacts, scope=scope, archive=None, redact=lambda text: text, run_id=RUN,
        unit=UNIT, worker="A_w2", events=[], max_attempts=3, recorded=lambda _id: 0,
    )  # fmt: skip
    ref = kept[0].version.ref
    state = deps.Store(worlds={WORLD: deps.WorldRecord(OWNER, {}, run_id=RUN)})
    state.events[WORLD] = [
        {"type": "artifact.created", "actor": "A_w2", "data": {"artifact": ref}},
        {"type": "artifact.delivered", "actor": "A_w2", "data": {"artifact": ref, "format": "in_app"}},
        {"type": "world.completed", "data": {"credits_used": 0}},
    ]
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides.update(
        {
            deps.get_store: lambda: state,
            deps.get_artifact_store: lambda: artifacts,
            result_store: lambda: lambda: artifacts,
            deps.get_current_principal: lambda: deps.Principal(OWNER, "owner", WORLD),
        }
    )
    return TestClient(app), writer, outcome


def served(client: TestClient, fmt: str) -> bytes:
    response = client.get(f"/worlds/{WORLD}/result")
    assert response.status_code == 200
    assert response.headers["content-type"].split(";")[0] == MIME[fmt]
    assert response.headers["content-disposition"] == f'attachment; filename="tilo-result.{fmt}"'
    return response.content


@pytest.fixture(autouse=True)
def in_process_render(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("TILO_RENDER_URL", raising=False)


def test_DEMO_DELIVER_pdf_opens_with_its_title_table_and_computed_column() -> None:
    client, writer, _ = deliver("pdf", json.dumps(SPEC))
    reader = PdfReader(io.BytesIO(served(client, "pdf")))
    text = " ".join(" ".join(page.extract_text() for page in reader.pages).split())  # table cells wrap
    assert len(reader.pages) >= 2 and "World time zones and cities" in text
    assert "Mumbai, Delhi, Kolkata" in text and "330" in text and "1999" not in text
    assert "DeliverableSpec" in writer.requests[0].prompt and len(writer.requests) == 1


def test_DEMO_DELIVER_docx_opens_as_a_word_document_with_headings_and_the_table() -> None:
    client, _, _ = deliver("docx", json.dumps(SPEC))
    with zipfile.ZipFile(io.BytesIO(served(client, "docx"))) as package:
        assert package.testzip() is None
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        document = ET.fromstring(package.read("word/document.xml"))
    words = list(document.itertext())
    assert {"Title", "Heading1"} <= {s.get(f"{ns}val") for s in document.iter(f"{ns}pStyle")}
    assert "World time zones and cities" in words and "Tokyo, Osaka" in words and "330" in words


def test_DEMO_DELIVER_pptx_opens_as_a_deck_of_slides_with_the_title_and_the_data() -> None:
    client, _, _ = deliver("pptx", json.dumps(SPEC))
    deck = Presentation(io.BytesIO(served(client, "pptx")))
    text = " ".join(
        run.text
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    )
    cells = " ".join(
        cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table for row in shape.table.rows
        for cell in row.cells
    )  # fmt: skip
    assert len(deck.slides) >= 3 and "World time zones and cities" in text and "Overview" in text
    assert "Mumbai, Delhi, Kolkata" in cells + text


def test_DEMO_DELIVER_xlsx_opens_as_a_workbook_with_typed_cells_formulas_and_totals() -> None:
    client, _, _ = deliver("xlsx", json.dumps(SPEC))
    body = served(client, "xlsx")
    formulas = load_workbook(io.BytesIO(body)).worksheets[-1]
    values = load_workbook(io.BytesIO(body), data_only=True).worksheets[-1]
    header = [cell.value for cell in values[1]]
    assert header[:6] == [*SPEC["tables"][0]["columns"], "Offset (minutes)"]
    assert values["B5"].value == 5.5 and values["F5"].value == 330 and formulas["F5"].value == "=B5*60"
    assert values["E3"].value == "=1+1 stays text" and formulas["E3"].data_type == "s"  # model text is never a formula


def test_DEMO_DELIVER_csv_is_a_real_csv_file_with_the_table_and_no_formula_cells() -> None:
    client, writer, outcome = deliver("csv", json.dumps(SPEC))
    body = served(client, "csv")
    assert outcome.outputs[0].media_type == "text/csv" and body.startswith(b"\xef\xbb\xbf")  # Excel reads UTF-8
    rows = list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))
    assert rows[0] == ["Time zone", "UTC offset (hours)", "Major cities", "Cities listed", "Note", "Offset (minutes)"]
    assert rows[1] == ["UTC-08:00 Pacific", "-8", "Los Angeles, Vancouver", "2", "West coast", "-480"]
    assert rows[2][4] == "'=1+1 stays text" and rows[4][1:2] + rows[4][5:] == ["5.5", "330"]
    assert len(rows) == 1 + len(ROWS) and b"\r\n" in body
    assert "A CSV file (csv) is one table" in writer.requests[0].prompt


def test_DEMO_DELIVER_csv_spec_without_a_table_gets_one_repair_then_the_file() -> None:
    no_table = {**SPEC, "tables": [], "charts": [], "sections": SPEC["sections"][:1]}
    client, writer, _ = deliver("csv", json.dumps(no_table), json.dumps(SPEC))
    assert len(writer.requests) == 2 and writer.requests[1].json_mode
    assert writer.requests[1].max_output_tokens > 4096  # a long table fits the repaired spec
    assert served(client, "csv").decode("utf-8-sig").startswith("Time zone,")


def test_DEMO_DELIVER_website_is_an_in_app_page_that_passes_the_checker_in_one_call() -> None:
    client, writer, outcome = deliver("website", PAGE)
    response = client.get(f"/worlds/{WORLD}/result")
    assert response.status_code == 200 and response.headers["content-type"] == "text/html; charset=utf-8"
    assert "sandbox allow-scripts" in response.headers["content-security-policy"]
    assert check_website(response.content).ok and b"window.scrollTo(0, 0);" in response.content
    assert b"<select" in response.content and b"<button" in response.content and len(writer.requests) == 1
    assert outcome.outputs[0].media_type == "text/html"
