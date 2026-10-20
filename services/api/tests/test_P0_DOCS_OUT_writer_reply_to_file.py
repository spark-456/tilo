"""P0_DOCS_OUT: one writer reply (a JSON spec, here from a fake provider) becomes a real PDF, Word or Excel file."""

import asyncio
import io
import json
import zipfile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from tilo_api import report_render
from tilo_harness.deliverable import DeliverableRefused
from tilo_harness.outputs import RenderedOffice, RenderedPdf

REPLY = json.dumps(
    {
        "title": "Event budget",
        "subtitle": "Autumn fair",
        "date": "1999-01-01",  # a model-written date is not a fact: the code's date wins
        "tables": [
            {
                "id": "t1",
                "columns": ["Item", "Qty", "Unit cost"],
                "rows": [["Chairs", 40, 3.5], ["Tables", 8, 20]],
                "formulas": [{"column": "Cost", "expr": "[Qty] * [Unit cost]"}],
            }
        ],
        "charts": [{"id": "c1", "kind": "bar", "table_id": "t1", "x": "Item", "y": "Cost", "title": "Cost by item"}],
        "sections": [
            {
                "heading": "Budget",
                "blocks": [
                    {"type": "paragraph", "text": "Totals below."},
                    {"type": "table", "table_id": "t1"},
                    {"type": "chart", "chart_id": "c1"},
                ],
            }
        ],
    }
)
MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def make(fmt: str, reply: str = REPLY):
    return asyncio.run(report_render.render_report(reply, fmt=fmt))


def test_P0_DOCS_OUT_pdf_reply_makes_a_pdf_with_cover_and_the_computed_column(monkeypatch):
    monkeypatch.delenv("TILO_RENDER_URL", raising=False)
    out = make("pdf")
    assert type(out) is RenderedPdf and out.media_type == "application/pdf" and out.content.startswith(b"%PDF-")
    pages = [p.extract_text() for p in PdfReader(io.BytesIO(out.content)).pages]
    assert "Event budget" in pages[0] and "Budget" not in pages[0].replace("Event budget", "")
    assert "Cost" in pages[1] and "140" in pages[1] and "160" in pages[1]


def test_P0_DOCS_OUT_docx_reply_makes_a_word_file_with_styled_headings(monkeypatch):
    monkeypatch.delenv("TILO_RENDER_URL", raising=False)
    out = make("docx")
    assert type(out) is RenderedOffice and out.media_type == MIME["docx"]
    with zipfile.ZipFile(io.BytesIO(out.content)) as package:
        assert package.testzip() is None
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        document = ET.fromstring(package.read("word/document.xml"))
        styles = {s.get(f"{ns}val") for s in document.iter(f"{ns}pStyle")}
        assert {"Title", "Heading1"} <= styles
        assert "140" in list(document.itertext()) and "styles.xml" in " ".join(package.namelist())
        assert len(document.findall(".//{*}chart")) == 1


def test_P0_DOCS_OUT_xlsx_reply_makes_an_excel_workbook_with_formulas_and_values(monkeypatch):
    monkeypatch.delenv("TILO_RENDER_URL", raising=False)
    out = make("xlsx")
    assert type(out) is RenderedOffice and out.media_type == MIME["xlsx"]
    formulas = load_workbook(io.BytesIO(out.content)).worksheets[-1]
    values = load_workbook(io.BytesIO(out.content), data_only=True).worksheets[-1]
    assert formulas["D2"].value == "=B2*C2" and values["D2"].value == 140 and values["D3"].value == 160
    assert formulas["D4"].value == "=SUM(D2:D3)" and values["D4"].value == 300


@pytest.mark.parametrize("fmt", ["pdf", "docx", "xlsx"])
def test_P0_DOCS_OUT_a_reply_with_a_formula_outside_arithmetic_is_refused_with_a_fixed_code(fmt):
    bad = json.loads(REPLY)
    bad["tables"][0]["formulas"] = [{"column": "Cost", "expr": "HYPERLINK([Qty])"}]
    with pytest.raises(DeliverableRefused) as refused:
        make(fmt, json.dumps(bad))
    assert refused.value.code == "invalid_spec" and "HYPERLINK" not in str(refused.value)
