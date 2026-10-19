"""DOCS-OUT A4: an uploaded workbook is found, shown to the writer as a layout, and filled from the writer's cells."""

import asyncio
import dataclasses
import io
import json
from datetime import UTC, datetime

import pytest
from _prod_wire_2b_lib import FakeBucket
from openpyxl import Workbook, load_workbook
from tilo_api import deps, report_render
from tilo_api.run_stores import GcsBlobBackend
from tilo_api.xlsx_template import attach_template, find_template
from tilo_harness.deliverable import TEMPLATE_ID, DeliverableRefused, parse_fill, prompt_name, render_format
from tilo_harness.executor import UnitContext
from tilo_harness.outputs import RenderedOffice
from tilo_harness.unit_prompt import UpstreamItem
from tilo_memory.artifacts import ArtifactStore, MemoryMetaBackend, NewVersion, Origin, Scope
from tilo_providers.safety import FakeSafetyProvider

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
RUN, WORLD, OWNER = "r_fill", "w_fill", "u_owner"
UNIT = {"id": "U9", "role": "writer", "unit_type": "deliver", "title": "Fill it"}
BRIEF = {"goal": "Fill my form", "deliverables": [{"title": "Form", "format": "xlsx"}]}


def workbook() -> bytes:
    wb = Workbook()
    wb.active.title = "Form"
    wb.active["A1"], wb.active["A2"] = "Name", "Amount"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


TEMPLATE = workbook()


def store_with(data: bytes, *, kind="upload", media=XLSX, run=RUN):
    artifacts = ArtifactStore(
        MemoryMetaBackend(),
        GcsBlobBackend(FakeBucket()),
        clock=lambda: datetime(2026, 10, 16, tzinfo=UTC),
        redactor=lambda text: text,
        archive_lookup=lambda *_: None,
        alert_sink=lambda _: None,
        screener=FakeSafetyProvider(),
    ).with_upload_text(lambda content, mime: "A1 Name A2 Amount")
    scope = Scope(deps.ORG_ID, WORLD, OWNER)
    origin = Origin("render", "U1", None) if kind == "render" else Origin(kind, None, None)
    item = artifacts.put(
        scope, NewVersion(media, data, run, origin, idempotency_key=None if kind == "upload" else f"{run}:U1:0:0")
    )
    stream = [{"type": "artifact.created", "data": {"artifact": item.ref}}]
    return artifacts, scope, stream


def test_DOCS_OUT_A4_the_uploaded_workbook_is_found_and_shown_to_the_writer_as_a_layout():
    artifacts, scope, stream = store_with(TEMPLATE)
    assert find_template(artifacts, scope, stream, RUN) == TEMPLATE
    context = UnitContext("A_w", (), brief=BRIEF)
    data, context = attach_template("xlsx", artifacts, scope, stream, RUN, context)
    assert data is not None
    first = context.upstream[0]
    assert first.unit_id == TEMPLATE_ID and "A1=Name" in first.text and first.tainted
    assert prompt_name(UNIT, context) == "unit_xlsx_fill.md"
    assert prompt_name(UNIT, dataclasses.replace(context, upstream=())) == "unit_pdf.md"
    assert prompt_name({**UNIT, "unit_type": "research"}, context) == "unit.md"


@pytest.mark.parametrize("case", ["wrong_format", "other_run", "not_an_upload", "not_xlsx", "unrecorded"])
def test_DOCS_OUT_A4_only_this_runs_uploaded_xlsx_counts_as_a_template(case):
    kwargs = {}
    if case == "other_run":
        kwargs["run"] = "r_else"
    if case == "not_an_upload":
        kwargs["kind"] = "render"
    if case == "not_xlsx":
        kwargs["media"] = "text/csv"
    data = b"a,b\n1,2\n" if case == "not_xlsx" else TEMPLATE
    artifacts, scope, stream = store_with(data, **kwargs)
    fmt = "pdf" if case == "wrong_format" else "xlsx"
    if case == "unrecorded":
        stream = []  # an artifact the world never announced is not a template
    context = UnitContext("A_w", (), brief=BRIEF)
    found, after = attach_template(fmt, artifacts, scope, stream, RUN, context)
    assert found is None and after.upstream == ()


def test_DOCS_OUT_A4_an_unreadable_workbook_falls_back_to_a_new_one():
    artifacts, scope, stream = store_with(b"PK\x03\x04 not a workbook")
    found, after = attach_template("xlsx", artifacts, scope, stream, RUN, UnitContext("A_w", (), brief=BRIEF))
    assert found is None and after.upstream == ()


GOOD = json.dumps(
    {"cells": [{"sheet": "Form", "cell": "B1", "value": "Ada"}, {"sheet": "Form", "cell": "B2", "value": 42}]}
)


def test_DOCS_OUT_A4_the_writers_cells_are_written_into_a_copy_of_the_template():
    output = asyncio.run(report_render.render_report(GOOD, fmt="xlsx", template=TEMPLATE))
    assert type(output) is RenderedOffice and output.media_type == XLSX
    sheet = load_workbook(io.BytesIO(output.content))["Form"]
    assert (sheet["A1"].value, sheet["B1"].value, sheet["B2"].value) == ("Name", "Ada", 42)
    assert json.loads(output.spec_text)["cells"][0]["value"] == "Ada"


@pytest.mark.parametrize(
    "reply",
    [
        "```json\n" + GOOD + "\n```",
        '{"cells": [], "extra": 1}',
        '{"cells": []}',
        '{"cells": [{"sheet": "Form", "cell": "B1", "value": null}]}',
        '{"cells": [{"sheet": "Form", "cell": "B1", "value": 1, "style": "x"}]}',
        '{"cells": [{"sheet": "Nope", "cell": "B1", "value": 1}]}',
        '{"cells": [{"sheet": "Form", "cell": "A1:B2", "value": 1}]}',
        '{"cells": [{"sheet": "Form", "cell": "B1", "value": 1}], "cells": []}',
        '{"cells": [{"sheet": "Form", "cell": "B1", "value": NaN}]}',
    ],
)
def test_DOCS_OUT_A4_a_flawed_fill_reply_makes_no_file(reply):
    with pytest.raises(DeliverableRefused):
        asyncio.run(report_render.render_report(reply, fmt="xlsx", template=TEMPLATE))


def test_DOCS_OUT_A4_parse_fill_and_the_format_choice_follow_the_brief():
    assert [c.cell for c in parse_fill(GOOD)] == ["B1", "B2"]
    assert render_format(UNIT, BRIEF) == "xlsx"


def test_DOCS_OUT_A4_without_a_template_the_reply_is_still_a_report_spec():
    spec = json.dumps(
        {"title": "T", "date": "2026-10-16", "tables": [{"id": "t", "columns": ["a"], "rows": [["x"]]}], "sections": []}
    )
    output = asyncio.run(report_render.render_report(spec, fmt="xlsx"))
    assert load_workbook(io.BytesIO(output.content)).worksheets[0]["A1"].value == "a"


def test_DOCS_OUT_A4_upstream_item_shape_matches_the_prompt_builder():
    item = UpstreamItem(TEMPLATE_ID, "Workbook to fill", "A1=Name", screening="clean", tainted=True)
    assert item.text == "A1=Name"


def test_DOCS_OUT_A4_the_writer_unit_sees_the_layout_and_its_cells_become_the_filled_file():
    from functools import partial

    from tilo_harness.tool_runner import ToolUnitRunner
    from tilo_providers.fake import FakeModelProvider
    from tilo_providers.interfaces import GenerateResponse

    class Script(FakeModelProvider):
        prompts: list[str] = []

        def generate(self, request):
            self.prompts.append(request.prompt)
            return GenerateResponse(text=GOOD, profile_id="script")

    provider = Script("script")
    artifacts, scope, stream = store_with(TEMPLATE)
    data, context = attach_template("xlsx", artifacts, scope, stream, RUN, UnitContext("A_w", (), brief=BRIEF))
    runner = ToolUnitRunner(
        {"writer": provider}, report_renderer=partial(report_render.render_report, fmt="xlsx", template=data)
    )
    outcome = asyncio.run(runner.run(UNIT, context))
    assert outcome.success and len(outcome.outputs) == 1
    assert outcome.outputs[0].media_type == XLSX
    assert "A1=Name" in provider.prompts[0] and "cells: a list of at most 2000" in provider.prompts[0]
    sheet = load_workbook(io.BytesIO(outcome.outputs[0].content))["Form"]
    assert (sheet["B1"].value, sheet["B2"].value) == ("Ada", 42)
