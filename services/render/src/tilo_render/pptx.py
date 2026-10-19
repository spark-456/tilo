"""PPTX deck renderer with native data-bound charts (FR-CON-09; PRD 6.6, 12.4, 12.6)."""

import base64
import io
import json
from pathlib import Path
from typing import Any

import pptx
import yaml
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
from tilo_contracts import Manifest

from tilo_render.charts import ChartData
from tilo_render.spec import (
    BulletsBlock,
    Chart,
    ChartBlock,
    ChartKind,
    DeliverableSpec,
    ParagraphBlock,
    SlideBlock,
    Table,
    TableBlock,
    slide_chunks,
)
from tilo_render.text import PROVENANCE_HEADING, provenance_lines, sanitize_text, wrap_long_token
from tilo_render.text import find_repo_root as _find_repo_root
from tilo_render.text import lines as _lines


def _resolve_color(role: str, tokens: dict[str, Any]) -> RGBColor:
    node: Any = tokens.get("color", {})
    for part in role.split("."):
        node = node[part]
    val = node["value"] if isinstance(node, dict) and "value" in node else str(node)
    return RGBColor.from_string(val.lstrip("#"))


def _chunk_text(text: str, max_chars: int = 1000) -> list[str]:
    if len(text) <= max_chars:
        return [text] if text else [""]
    chunks: list[str] = []
    while text:
        if len(text) <= max_chars:
            chunks.append(text)
            break
        sp = text.rfind(" ", 0, max_chars)
        idx = max_chars if sp <= 0 else sp + 1
        chunks.append(text[:idx])
        text = text[idx:]
    return chunks or [""]


def _slide(
    prs: Any, bg: RGBColor, heading: str | None = None, h_col: RGBColor | None = None, h_pt: float = 24.0
) -> Any:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = bg
    if heading and h_col:
        tb = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.0))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.text, p.font.size, p.font.bold, p.font.color.rgb = sanitize_text(heading), Pt(h_pt), True, h_col
    return slide


def _body_tf(slide: Any) -> Any:
    tf = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(4.5)).text_frame
    tf.word_wrap = True
    return tf


def _add_title_slide(prs: Any, spec: DeliverableSpec, col: dict[str, RGBColor], typ: dict[str, float]) -> None:
    slide = _slide(prs, col["title_background"])
    tf = slide.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.3), Inches(3.5)).text_frame
    p0 = tf.paragraphs[0]
    p0.text, p0.font.size, p0.font.bold, p0.font.color.rgb = (
        sanitize_text(spec.title),
        Pt(typ["title_size_pt"]),
        True,
        col["title_text"],
    )
    if spec.subtitle:
        p1 = tf.add_paragraph()
        p1.text, p1.font.size, p1.font.color.rgb = (
            sanitize_text(spec.subtitle),
            Pt(typ["subtitle_size_pt"]),
            col["title_subtext"],
        )
    p2 = tf.add_paragraph()
    p2.text, p2.font.size, p2.font.color.rgb = sanitize_text(spec.date), Pt(14), col["title_subtext"]


def _add_para_slide(prs: Any, heading: str, text: str, col: dict[str, RGBColor], typ: dict[str, float]) -> None:
    tf = _body_tf(_slide(prs, col["slide_background"], heading, col["slide_header"], typ["slide_title_size_pt"]))
    # One paragraph per line: python-pptx would turn "\n" in `p.text` into a line break inside ONE paragraph.
    # A blank line stays an (empty) paragraph so the visual paragraph break survives.
    for idx, line in enumerate(_lines(text)):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text, p.font.size, p.font.color.rgb = line, Pt(typ["body_size_pt"]), col["slide_text"]


def _add_bullets_slide(
    prs: Any, heading: str, bullets: list[str], col: dict[str, RGBColor], typ: dict[str, float]
) -> None:
    tf = _body_tf(_slide(prs, col["slide_background"], heading, col["slide_header"], typ["slide_title_size_pt"]))
    # Every non-blank line of every item is its own bullet paragraph with its own marker; blank lines are
    # paragraph breaks, never empty bullets.
    lines = [line for item in bullets for line in _lines(item) if line.strip()]
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"•  {line}"
        p.level = 0
        p.font.size, p.font.color.rgb = Pt(typ["bullet_size_pt"]), col["slide_text"]


def _add_table_slide(
    prs: Any, heading: str, table: Table, chunk: list[list[Any]], col: dict[str, RGBColor], typ: dict[str, float]
) -> None:
    slide = _slide(prs, col["slide_background"], heading, col["slide_header"], typ["slide_title_size_pt"])
    tbl = slide.shapes.add_table(
        len(chunk) + 1, len(table.columns), Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.5)
    ).table
    for c_idx, c_name in enumerate(table.columns):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = col["table_header_bg"]
        p = cell.text_frame.paragraphs[0]
        p.text, p.font.size, p.font.bold, p.font.color.rgb = (
            sanitize_text(c_name),
            Pt(typ["table_header_size_pt"]),
            True,
            col["table_header_fg"],
        )
    for r_idx, row in enumerate(chunk):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = col["table_row_bg"]
            p = cell.text_frame.paragraphs[0]
            p.text, p.font.size, p.font.color.rgb = (
                sanitize_text(str(val)),
                Pt(typ["table_body_size_pt"]),
                col["table_row_fg"],
            )


def _add_chart_slide(
    prs: Any, spec: DeliverableSpec, heading: str, chart_def: Chart, col: dict[str, RGBColor], typ: dict[str, float]
) -> None:
    slide = _slide(prs, col["slide_background"], heading, col["slide_header"], typ["slide_title_size_pt"])
    cdata = CategoryChartData()
    cd = ChartData.from_spec(spec, chart_def.id)
    cdata.categories = [wrap_long_token(sanitize_text(str(c))) for c in cd.categories]
    seen_names: set[str] = set()
    for s_name, s_vals in cd.series.items():
        name = sanitize_text(str(s_name))
        while name in seen_names:
            name = f"{name}_"
        seen_names.add(name)
        cdata.add_series(name, tuple(s_vals))
    kinds = {
        ChartKind.BAR: XL_CHART_TYPE.COLUMN_CLUSTERED,
        ChartKind.LINE: XL_CHART_TYPE.LINE,
        ChartKind.PIE: XL_CHART_TYPE.PIE,
    }
    chart = slide.shapes.add_chart(
        kinds[chart_def.kind], Inches(1.5), Inches(2.0), Inches(10.3), Inches(4.8), cdata
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = sanitize_text(chart_def.title)
    if chart.chart_title.text_frame.paragraphs:
        chart.chart_title.text_frame.paragraphs[0].font.size = Pt(typ["chart_title_size_pt"])
        chart.chart_title.text_frame.paragraphs[0].font.color.rgb = col["slide_header"]


def _add_design(prs: Any, block: SlideBlock, col: dict[str, RGBColor], typ: dict[str, float]) -> None:
    left, right = slide_chunks(block.left), slide_chunks(block.right)
    for index in range(max(len(left), len(right))):
        slide = _slide(
            prs,
            col["slide_background"],
            "\n".join(_chunk_text(block.title, 45)),
            col["slide_header"],
            min(16, typ["slide_title_size_pt"]),
        )
        slide.notes_slide.notes_text_frame.text = sanitize_text(block.notes)
        for column, chunks in enumerate((left, right)):
            text = chunks[index] if index < len(chunks) else ""
            x, width = (1.0, 5.3) if column == 0 else (7.0, 5.3)
            y, height = 2.2, 4.2
            size = typ["body_size_pt"]
            if block.layout == "stat" and column == 0 and len(text) <= 10 and "\n" not in text:
                size = 44
            if block.layout == "image" and column == 1:
                continue
            tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height)).text_frame
            tf.word_wrap = True
            for number, line in enumerate(text.split("\n")):
                para = tf.paragraphs[0] if number == 0 else tf.add_paragraph()
                para.text, para.font.size, para.font.color.rgb = sanitize_text(line), Pt(size), col["slide_text"]
            if block.layout == "image" and block.image_base64:
                picture = slide.shapes.add_picture(
                    io.BytesIO(base64.b64decode(block.image_base64)), Inches(7), Inches(2.2)
                )
                scale = min(Inches(5.3) / picture.width, Inches(4.2) / picture.height)
                picture.width, picture.height = int(picture.width * scale), int(picture.height * scale)


def render_pptx(
    spec: DeliverableSpec,
    *,
    manifest: Manifest | None = None,
    template_path: Path | str | None = None,
    tokens_path: Path | str | None = None,
) -> bytes:
    root = _find_repo_root()
    tmpl_p = Path(template_path) if template_path else root / "templates" / "deck" / "default.yaml"
    tok_p = Path(tokens_path) if tokens_path else root / "design" / "tokens.json"

    with open(tmpl_p, encoding="utf-8") as f:
        tmpl = yaml.safe_load(f)
    with open(tok_p, encoding="utf-8") as f:
        tokens = json.load(f)

    colors = {k: _resolve_color(v, tokens) for k, v in tmpl.get("colors", {}).items()}
    typo = {k: float(v) for k, v in tmpl.get("typography", {}).items()}
    rows_per_slide = int(tmpl.get("table", {}).get("rows_per_slide", 10))
    bullets_per_slide = int(tmpl.get("bullets", {}).get("items_per_slide", 8))
    chars_per_slide = int(tmpl.get("paragraph", {}).get("chars_per_slide", 1000))

    prs = pptx.Presentation()
    prs.core_properties.title, prs.core_properties.author = sanitize_text(spec.title), "Tilo"
    size = tmpl.get("slide_size", {})
    prs.slide_width, prs.slide_height = (
        Inches(float(size.get("width_in", 13.333))),
        Inches(float(size.get("height_in", 7.5))),
    )

    _add_title_slide(prs, spec, colors, typo)

    for sec in spec.sections:
        _slide(prs, colors["section_background"], sec.heading, colors["section_title"], typo["section_title_size_pt"])
        for b in sec.blocks:
            if isinstance(b, SlideBlock):
                _add_design(prs, b, colors, typo)
            elif isinstance(b, ParagraphBlock):
                for chunk in _chunk_text(b.text, max_chars=chars_per_slide):
                    _add_para_slide(prs, sec.heading, chunk, colors, typo)
            elif isinstance(b, BulletsBlock):
                items = b.bullets or [""]
                for i in range(0, len(items), bullets_per_slide):
                    _add_bullets_slide(prs, sec.heading, items[i : i + bullets_per_slide], colors, typo)
            elif isinstance(b, TableBlock):
                t = next(tbl for tbl in spec.tables if tbl.id == b.table_id)
                for i in range(0, max(1, len(t.rows)), rows_per_slide):
                    _add_table_slide(prs, sec.heading, t, t.rows[i : i + rows_per_slide], colors, typo)
            elif isinstance(b, ChartBlock):
                c = next(ch for ch in spec.charts if ch.id == b.chart_id)
                _add_chart_slide(prs, spec, sec.heading, c, colors, typo)

    if manifest is not None:
        prov = provenance_lines(manifest)
        for i in range(0, len(prov), bullets_per_slide):
            _add_bullets_slide(prs, PROVENANCE_HEADING, prov[i : i + bullets_per_slide], colors, typo)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
