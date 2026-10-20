"""Excel workbooks from the tables of a DeliverableSpec (DOCS-OUT A4; `xlsxwriter` writes, `openpyxl` fills a template).

Each table becomes a sheet: a branded header row, frozen and filtered, typed numbers and ISO dates, column widths from
the content, computed columns as real row formulas (`=B2*C2`), a totals row of same-sheet `SUM` formulas under the
numeric columns (every formula also carries its computed value, so a viewer that does not recalculate still shows it),
and a native Excel chart for each chart block. Everything a model wrote is
data: a string is always stored as a string (never a formula, link or macro), and the workbook has no external parts.
Total cells are capped so the file stays small and can be screened.
"""

import io
import re
import zipfile
from collections.abc import Sequence
from datetime import date, datetime
from math import isfinite
from typing import Any

import xlsxwriter
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter
from pydantic import BaseModel, ConfigDict, Field, StrictFloat, StrictInt, StrictStr, field_validator

from tilo_render import formula
from tilo_render.spec import BulletsBlock, Chart, DeliverableSpec, ParagraphBlock, Table, TableBlock
from tilo_render.text import brand_color, sanitize_text

MAX_XLSX_CELLS = 7000  # table cells (header and totals included) plus report lines, across all sheets
MAX_CELL_CHARS = 32_000  # Excel's own limit is 32,767
_ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
_NOT_A_QUANTITY = re.compile(r"\b(year|id|rank|code|zip|phone|index|no)\b|#", re.IGNORECASE)
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_FILL_CELLS = 2000  # cells one writer reply may fill in a template
MAX_FILL_TEXT = 2000
MAX_TEMPLATE_BYTES = 8 * 1024 * 1024
MAX_TEMPLATE_EXPANDED = 32 * 1024 * 1024
_ADDRESS = re.compile(r"([A-Z]{1,3})([1-9][0-9]{0,4})")
_ACTIVE_PARTS = ("externallink", "vbaproject", "activex", "embeddings/", "oleobject")


def text_value(value: Any) -> str:
    return sanitize_text(str(value))[:MAX_CELL_CHARS]


def store(cell: Any, value: Any) -> None:
    """One cell: numbers stay numbers, ISO dates become dates, every other string is stored as a string."""
    if isinstance(value, bool) or not isinstance(value, (int, float, str)):
        value = str(value)
    if isinstance(value, (int, float)) and isfinite(value):
        cell.value = value
        return
    plain = text_value(value)
    if _ISO_DATE.fullmatch(plain):
        try:
            cell.value = date.fromisoformat(plain)
            if cell.number_format == "General":  # a template's own date format wins
                cell.number_format = "yyyy-mm-dd"
            return
        except ValueError:
            pass
    cell.value = plain
    cell.data_type = "s"  # a leading "=", "+", "-" or "@" is text here, never a formula


def _sheet_name(base: str, used: set[str]) -> str:
    name = _BAD_SHEET_CHARS.sub(" ", sanitize_text(base)).replace("\n", " ").strip().strip("'").strip() or "Sheet"
    candidate, n = name[:31], 1
    while candidate.lower() in used:
        n += 1
        suffix = f" {n}"
        candidate = name[: 31 - len(suffix)].rstrip().strip("'") + suffix
    used.add(candidate.lower())
    return candidate


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and isfinite(value)


class _Styles:
    def __init__(self, book: Any) -> None:
        brand = brand_color()
        head = {"bold": True, "font_color": "#FFFFFF", "bg_color": brand, "valign": "vcenter", "border": 1}
        self.header = book.add_format({**head, "border_color": "#D9DEE5", "text_wrap": True})
        self.text = book.add_format({"valign": "top", "text_wrap": True})
        self.integer = book.add_format({"num_format": "#,##0", "valign": "top"})
        self.decimal = book.add_format({"num_format": "#,##0.00", "valign": "top"})
        self.date = book.add_format({"num_format": "yyyy-mm-dd", "align": "left", "valign": "top"})
        self.total_label = book.add_format({"bold": True, "top": 1})
        self.total_int = book.add_format({"bold": True, "top": 1, "num_format": "#,##0"})
        self.total_dec = book.add_format({"bold": True, "top": 1, "num_format": "#,##0.00"})
        self.title = book.add_format({"bold": True, "font_size": 16, "text_wrap": True, "valign": "top"})
        self.heading = book.add_format({"bold": True, "font_size": 12, "font_color": brand, "valign": "top"})
        self.muted = book.add_format({"font_color": "#5A6573", "text_wrap": True, "valign": "top"})


def _write(sheet: Any, styles: _Styles, row: int, column: int, value: Any) -> None:
    """One cell (zero-based): numbers stay numbers, ISO dates become dates, every other string is a string."""
    if _is_number(value):
        sheet.write_number(row, column, value, styles.integer if isinstance(value, int) else styles.decimal)
        return
    plain = text_value(value)
    if _ISO_DATE.fullmatch(plain):
        try:
            day = datetime.combine(date.fromisoformat(plain), datetime.min.time())
            sheet.write_datetime(row, column, day, styles.date)
            return
        except ValueError:
            pass
    sheet.write_string(row, column, plain, styles.text)  # a leading "=", "+", "-" or "@" is text, never a formula


def _table_sheet(sheet: Any, styles: _Styles, table: Table) -> int:
    """Fills one sheet from a table: header, typed cells, formula columns and a totals row. Returns the last row."""
    sheet.set_row(0, 22)
    for column, name in enumerate(table.columns):
        sheet.write_string(0, column, text_value(name), styles.header)
    letters = {name: get_column_letter(c) for c, name in enumerate(table.columns, 1)}
    computed = {f.column: formula.parse(f.expr) for f in table.formulas}
    for r, row in enumerate(table.rows, 1):
        for c, value in enumerate(row):
            name = table.columns[c]
            if name in computed and _is_number(value):
                refs = {n: f"{letters[n]}{r + 1}" for n in formula.columns_used(computed[name])}
                style = styles.integer if isinstance(value, int) else styles.decimal
                sheet.write_formula(r, c, "=" + formula.to_excel(computed[name], refs), style, value)
            else:
                _write(sheet, styles, r, c, value)
    last = len(table.rows) + 1
    summed = [
        c
        for c, name in enumerate(table.columns)
        if all(_is_number(row[c]) for row in table.rows) and not _NOT_A_QUANTITY.search(name)
    ]
    if summed:
        if 0 not in summed:
            sheet.write_string(last, 0, "Total", styles.total_label)
        for c in summed:
            letter = get_column_letter(c + 1)
            whole = all(isinstance(row[c], int) for row in table.rows)
            style = styles.total_int if whole else styles.total_dec
            sheet.write_formula(last, c, f"=SUM({letter}2:{letter}{last})", style, sum(row[c] for row in table.rows))
    for c, name in enumerate(table.columns):
        longest = max([len(text_value(name)), *(len(text_value(row[c])) for row in table.rows)])
        sheet.set_column(c, c, min(max(10, longest + 2), 60))
    sheet.freeze_panes(1, 0)
    sheet.autofilter(0, 0, last - 1, len(table.columns) - 1)
    return last


_CHART_TYPES = {"bar": "column", "line": "line", "pie": "pie"}


def _add_chart(book: Any, sheet: Any, table: Table, chart: Chart, slot: int) -> None:
    """A native Excel chart that reads the table's own cells, so it follows any edit to them."""
    native = book.add_chart({"type": _CHART_TYPES[chart.kind.value]})
    name = sheet.get_name()
    x = table.columns.index(chart.x)
    for y_name in chart.y_columns:
        y = table.columns.index(y_name)
        native.add_series(
            {
                "name": [name, 0, y],
                "categories": [name, 1, x, len(table.rows), x],
                "values": [name, 1, y, len(table.rows), y],
            }
        )
    native.set_title({"name": text_value(chart.title)})
    native.set_size({"width": 560, "height": 320})
    sheet.insert_chart(1 + slot * 17, len(table.columns) + 1, native)


def _cell_count(spec: DeliverableSpec) -> int:
    tables = sum((len(t.rows) + 2) * len(t.columns) for t in spec.tables)
    lines = sum(len(b.bullets) if isinstance(b, BulletsBlock) else 1 for s in spec.sections for b in s.blocks)
    return tables + lines + 3 + len(spec.sections)


def _report_sheet(sheet: Any, styles: _Styles, spec: DeliverableSpec) -> None:
    sheet.set_column(0, 0, 100)
    row = 0
    for text, style in ((spec.title, styles.title), (spec.subtitle or "", styles.muted), (spec.date, styles.muted)):
        if text:
            sheet.write_string(row, 0, text_value(text), style)
            row += 1
    for section in spec.sections:
        row += 1
        sheet.write_string(row, 0, text_value(section.heading), styles.heading)
        for block in section.blocks:
            if isinstance(block, ParagraphBlock):
                items = [block.text]
            elif isinstance(block, BulletsBlock):
                items = [f"- {item}" for item in block.bullets]
            else:
                continue
            for item in items:
                row += 1
                sheet.write_string(row, 0, text_value(item), styles.text)


def _fixed_zip(data: bytes) -> bytes:
    """The same package with a fixed timestamp on every part, so one spec always gives the same bytes."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist(), key=lambda n: (n != "[Content_Types].xml", n)):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            target.writestr(info, source.read(name), compress_type=zipfile.ZIP_DEFLATED)
    return out.getvalue()


def render_xlsx(spec: DeliverableSpec) -> bytes:
    # Revalidate objects constructed without validation before creating any workbook parts.
    spec = DeliverableSpec.model_validate(spec.model_dump())
    if _cell_count(spec) > MAX_XLSX_CELLS:
        raise ValueError("too many cells for one workbook")
    buffer = io.BytesIO()
    book = xlsxwriter.Workbook(
        buffer,
        {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False, "strings_to_numbers": False},
    )
    book.set_properties(
        {"author": "Tilo", "title": text_value(spec.title)[:200], "created": datetime.fromisoformat(spec.date)}
    )
    styles = _Styles(book)
    used: set[str] = set()
    has_text = any(isinstance(b, ParagraphBlock | BulletsBlock) for s in spec.sections for b in s.blocks)
    if has_text or not spec.tables:
        _report_sheet(book.add_worksheet(_sheet_name("Summary" if spec.tables else "Report", used)), styles, spec)
    headings: dict[str, str] = {}  # the heading of the section that first shows each table
    for section in spec.sections:
        for block in section.blocks:
            if isinstance(block, TableBlock):
                headings.setdefault(block.table_id, section.heading)
    sheets: dict[str, Any] = {}
    for table in spec.tables:
        sheets[table.id] = book.add_worksheet(_sheet_name(headings.get(table.id, table.id), used))
        _table_sheet(sheets[table.id], styles, table)
    placed: dict[str, int] = {}
    for chart in spec.charts:
        table = next(t for t in spec.tables if t.id == chart.table_id)
        _add_chart(book, sheets[table.id], table, chart, placed.get(table.id, 0))
        placed[table.id] = placed.get(table.id, 0) + 1
    book.close()
    return _fixed_zip(buffer.getvalue())


class FillCell(BaseModel):
    """One value the writer chose for one cell of the user's template. Data only: never a formula."""

    model_config = ConfigDict(extra="forbid", frozen=True)
    sheet: str = Field(min_length=1, max_length=31)
    cell: str
    value: StrictStr | StrictInt | StrictFloat

    @field_validator("cell")
    @classmethod
    def _plain_address(cls, value: str) -> str:
        found = _ADDRESS.fullmatch(value)
        if found is None or column_index_from_string(found.group(1)) > 16384:
            raise ValueError("a cell is one plain address such as B7")
        return value

    @field_validator("value")
    @classmethod
    def _bounded(cls, value: str | int | float) -> str | int | float:
        if isinstance(value, float) and not isfinite(value):
            raise ValueError("a number must be finite")
        if isinstance(value, str) and len(value) > MAX_FILL_TEXT:
            raise ValueError("text is too long")
        return value


def _open_template(data: bytes) -> Workbook:
    """The template as a workbook, or ValueError: bounded, with no macros, embedded objects or external links."""
    if not isinstance(data, bytes) or not 4 < len(data) <= MAX_TEMPLATE_BYTES:
        raise ValueError("the template is not a bounded workbook")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            infos = package.infolist()
            names = [info.filename.lower() for info in infos]
            if len(infos) > 512 or sum(info.file_size for info in infos) > MAX_TEMPLATE_EXPANDED:
                raise ValueError("the template is too large")
            if any(marker in name for name in names for marker in _ACTIVE_PARTS) or any(
                name.endswith(".bin") and "printersettings" not in name for name in names
            ):
                raise ValueError("the template holds active or linked content")
        return load_workbook(io.BytesIO(data), keep_links=False, keep_vba=False)
    except ValueError:
        raise
    except Exception:
        raise ValueError("the template is not a readable workbook") from None


def fill_workbook(template: bytes, fills: Sequence[FillCell]) -> bytes:
    """A copy of the user's workbook with the writer's values in the named cells. Formatting, widths, merges, number
    formats and the template's own formulas stay. A fill may not create a sheet, write into a merged range other than
    its first cell, or overwrite a formula; any such fill refuses the whole reply (no half-filled file)."""
    if len(fills) > MAX_FILL_CELLS:
        raise ValueError("too many cells to fill")
    workbook = _open_template(template)
    for fill in (FillCell.model_validate(f.model_dump()) for f in fills):
        if fill.sheet not in workbook.sheetnames:
            raise ValueError("no such sheet in the template")
        cell = workbook[fill.sheet][fill.cell]
        if isinstance(cell, MergedCell) or cell.data_type == "f":
            raise ValueError("that cell is merged or holds a formula")
        store(cell, fill.value)
    workbook.calculation.fullCalcOnLoad = True  # the template's own formulas recompute when it is opened
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def describe_workbook(template: bytes, *, max_chars: int = 3500) -> str:
    """The template's sheets, merged ranges and filled cells as `A1=Label` (formulas as `[formula]`): what a writer
    needs to choose where values go. Bounded in rows, columns and characters."""
    workbook = _open_template(template)
    out: list[str] = []
    for sheet in workbook.worksheets[:10]:
        merged = ", ".join(str(r) for r in list(sheet.merged_cells.ranges)[:10])
        out.append(f'Sheet "{text_value(sheet.title)}" ({sheet.dimensions}){f"; merged {merged}" if merged else ""}')
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 200), max_col=min(sheet.max_column, 40)):
            items = []
            for cell in row:
                if cell.value is None:
                    continue
                shown = "[formula]" if cell.data_type == "f" else " ".join(text_value(cell.value)[:60].split())
                items.append(f"{cell.coordinate}={shown}")
            if items:
                out.append(" | ".join(items))
    text = "\n".join(out)
    return text if len(text) <= max_chars else text[: max_chars - 30].rstrip() + "\n(more cells left out)"
