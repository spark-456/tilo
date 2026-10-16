import base64
import io

import pptx
import pytest
from pydantic import ValidationError
from tilo_render.pptx import render_pptx
from tilo_render.spec import DeliverableSpec, RequestDeliverableSpec


def _spec(blocks):
    return dict(title="Evidence", date="2026-10-16", sections=[dict(heading="Findings", blocks=blocks)])


def test_GEN_06_structured_layouts_notes_and_overflow():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (40, 30), "blue").save(buf, format="PNG")
    blocks = [
        dict(
            type="slide",
            layout="two_column",
            title="Compare",
            left="Left evidence",
            right="Right evidence",
            notes="Explain comparison",
        ),
        dict(type="slide", layout="stat", title="Measured", left="42%", right="Recorded rate"),
        dict(
            type="slide",
            layout="image",
            title="Source image",
            left="Figure caption",
            image_base64=base64.b64encode(buf.getvalue()).decode(),
        ),
        dict(type="slide", layout="two_column", title="Long", left="word " * 600, right="Other evidence"),
    ]
    deck = pptx.Presentation(io.BytesIO(render_pptx(DeliverableSpec.model_validate(_spec(blocks)))))
    assert len(deck.slides) > 6
    assert deck.slides[2].notes_slide.notes_text_frame.text == "Explain comparison"
    assert any(shape.shape_type == 13 for shape in deck.slides[4].shapes)
    assert all(
        shape.left >= 0
        and shape.top >= 0
        and shape.left + shape.width <= deck.slide_width
        and shape.top + shape.height <= deck.slide_height
        for slide in deck.slides
        for shape in slide.shapes
    )
    texts = [shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]
    assert sum(text.count("word") for text in texts) == 600


def test_GEN_06_slide_work_includes_notes_and_continuations():
    with pytest.raises(ValidationError, match="character budget"):
        RequestDeliverableSpec.model_validate(
            _spec([dict(type="slide", layout="stat", title="Metric", left="1", notes="x" * 20000) for _ in range(3)])
        )
    with pytest.raises(ValidationError):
        DeliverableSpec.model_validate(
            _spec([dict(type="slide", layout="image", title="Bad", image_base64="file:///private")])
        )
