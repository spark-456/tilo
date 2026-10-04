"""Acceptance tests for AGR-04, provenance appendix in PPTX and PDF deliverables (FR-CON-10)."""

import io
from typing import Any

import pptx
from pypdf import PdfReader
from tilo_contracts import Manifest, ManifestProducer, ManifestSource
from tilo_render.pdf import render_pdf
from tilo_render.pptx import render_pptx
from tilo_render.spec import DeliverableSpec, ParagraphBlock, Section
from tilo_render.text import find_repo_root


def _spec() -> DeliverableSpec:
    return DeliverableSpec(
        title="Quarterly Review",
        date="2026-10-02",
        sections=[Section(heading="Summary", blocks=[ParagraphBlock(text="Key results.")])],
    )


def _valid_manifests() -> list[Manifest]:
    fixture_dir = find_repo_root() / "contracts" / "fixtures" / "manifest" / "valid"
    manifests = [
        Manifest.model_validate_json(p.read_text(encoding="utf-8")) for p in sorted(fixture_dir.glob("*.json"))
    ]
    assert len(manifests) >= 2
    return manifests


def _pptx_slide_text(slide: Any) -> str:
    return "\n".join(sh.text for sh in slide.shapes if sh.has_text_frame)


def _pdf_text(raw: bytes) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(raw)).pages)


def _assert_manifest_in_text(manifest: Manifest, text: str) -> None:
    assert "Sources and provenance" in text
    assert manifest.produced_by.agent_id in text
    assert manifest.produced_by.role in text
    if manifest.produced_by.skill is not None:
        assert manifest.produced_by.skill.id in text
        assert manifest.produced_by.skill.version in text
    else:
        assert "Skill: none" in text
    if manifest.sources:
        for src in manifest.sources:
            assert src.id in text
            assert src.ref in text
    else:
        assert "Sources: none" in text
    if manifest.claims and all(c.verified for c in manifest.claims):
        assert "Verifier: verified" in text
    else:
        assert "Verifier: unverified" in text
    if manifest.authz is not None:
        assert manifest.authz.ref in text
        assert manifest.authz.by in text
        assert manifest.authz.at in text
    else:
        assert "Approval: none" in text


def test_FR_CON_10_pptx_appendix_lists_every_source_skill_agent_and_approval() -> None:
    spec = _spec()
    for manifest in _valid_manifests():
        prs = pptx.Presentation(io.BytesIO(render_pptx(spec, manifest=manifest)))
        _assert_manifest_in_text(manifest, _pptx_slide_text(prs.slides[-1]))


def test_FR_CON_10_pdf_appendix_lists_every_source_skill_agent_and_approval() -> None:
    spec = _spec()
    for manifest in _valid_manifests():
        reader = PdfReader(io.BytesIO(render_pdf(spec, manifest=manifest)))
        _assert_manifest_in_text(manifest, reader.pages[-1].extract_text())


def test_FR_CON_10_no_manifest_adds_no_appendix() -> None:
    spec = _spec()
    manifest = _valid_manifests()[0]

    prs_without = pptx.Presentation(io.BytesIO(render_pptx(spec)))
    prs_none = pptx.Presentation(io.BytesIO(render_pptx(spec, manifest=None)))
    prs_with = pptx.Presentation(io.BytesIO(render_pptx(spec, manifest=manifest)))
    assert len(prs_without.slides) == len(prs_none.slides) == len(prs_with.slides) - 1
    assert all("Sources and provenance" not in _pptx_slide_text(s) for s in prs_without.slides)

    pdf_without = _pdf_text(render_pdf(spec))
    pdf_none = _pdf_text(render_pdf(spec, manifest=None))
    assert "Sources and provenance" not in pdf_without
    assert "Sources and provenance" not in pdf_none


def test_FR_CON_10_source_refs_are_text_not_links() -> None:
    spec = _spec()
    base = _valid_manifests()[0]
    evil_ref = 'https://example.com/report <a href="https://evil.invalid">click</a> & \ufffe\uffff\x00\x07'
    untrusted = base.model_copy(
        update={
            "produced_by": ManifestProducer.model_construct(
                agent_id="A_writer1\ufffe\x00",
                role="writer",
                skill=base.produced_by.skill,
                model_role=base.produced_by.model_role,
                profile=base.produced_by.profile,
            ),
            "sources": [
                *base.sources,
                ManifestSource.model_construct(
                    id="src_untrusted",
                    kind="web",
                    ref=evil_ref,
                    fingerprint=base.sources[0].fingerprint,
                    fetched_at=base.sources[0].fetched_at,
                ),
            ],
        }
    )

    prs = pptx.Presentation(io.BytesIO(render_pptx(spec, manifest=untrusted)))
    last_text = _pptx_slide_text(prs.slides[-1])
    assert "https://example.com/quarterly-report" in last_text
    assert '<a href="https://evil.invalid">click</a>' in last_text
    for bad in ("\ufffe", "\uffff", "\x00", "\x07"):
        assert bad not in last_text
    for slide in prs.slides:
        assert "hlinkClick" not in slide._element.xml
        assert all("hyperlink" not in rel.reltype for rel in slide.part.rels.values())
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        assert run.hyperlink.address is None

    raw_pdf = render_pdf(spec, manifest=untrusted)
    reader = PdfReader(io.BytesIO(raw_pdf))
    pdf_text = "\n".join(p.extract_text() for p in reader.pages)
    assert "https://example.com/quarterly-report" in pdf_text
    assert '<a href="https://evil.invalid">click</a>' in pdf_text
    for bad in ("\ufffe", "\uffff", "\x00", "\x07"):
        assert bad not in pdf_text
    for page in reader.pages:
        assert not page.get("/Annots")
    assert b"/URI" not in raw_pdf and b"/Link" not in raw_pdf
