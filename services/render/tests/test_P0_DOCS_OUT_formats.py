"""P0_DOCS_OUT: PDF, Word and Excel deliverables are real, well-formatted files made by code from one JSON spec."""

import io
import zipfile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from pypdf import PdfReader
from tilo_render.docx import render_docx
from tilo_render.pdf import render_pdf
from tilo_render.spec import DeliverableSpec
from tilo_render.xlsx import render_xlsx

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def spec(**extra) -> DeliverableSpec:
    body = {
        "title": "Q3 Sales Review",
        "subtitle": "North and South regions",
        "date": "2026-10-20",
        "tables": [
            {
                "id": "t1",
                "columns": ["Region", "Units", "Unit price"],
                "rows": [["North", 3, 10.5], ["South", 4, 20.25], ["East", 5, 8]],
                "formulas": [{"column": "Revenue", "expr": "[Units] * [Unit price]"}],
            }
        ],
        "charts": [{"id": "c1", "kind": "bar", "table_id": "t1", "x": "Region", "y": "Revenue", "title": "Revenue"}],
        "sections": [
            {
                "heading": "Summary",
                "blocks": [
                    {"type": "paragraph", "text": "Revenue is up."},
                    {"type": "bullets", "bullets": ["North grew", "South led"]},
                    {"type": "table", "table_id": "t1"},
                    {"type": "chart", "chart_id": "c1"},
                ],
            },
            {"heading": "Next steps", "blocks": [{"type": "paragraph", "text": "Review in Q4."}]},
        ],
    }
    body.update(extra)
    return DeliverableSpec.model_validate(body)


# ---- the spec: formulas are computed columns, deterministic, arithmetic only --------------------------------------


def test_P0_DOCS_OUT_formula_column_is_appended_with_computed_values_and_survives_a_round_trip():
    s = spec()
    table = s.tables[0]
    assert table.columns == ["Region", "Units", "Unit price", "Revenue"]
    assert [row[3] for row in table.rows] == [31.5, 81.0, 40]
    again = DeliverableSpec.model_validate(s.model_dump())
    assert again.tables[0].columns == table.columns and again.tables[0].rows == table.rows
    assert again == s


@pytest.mark.parametrize(
    "expr",
    ["[Nope] + 1", "HYPERLINK([Units])", "[Units] ** 2", "[Units] / 0", "__import__('os')", "[Units] +", "A1 + 2", ""],
)
def test_P0_DOCS_OUT_formula_outside_the_arithmetic_grammar_is_refused(expr):
    with pytest.raises(ValidationError):
        spec(
            tables=[
                {
                    "id": "t1",
                    "columns": ["Region", "Units"],
                    "rows": [["N", 3]],
                    "formulas": [{"column": "X", "expr": expr}],
                }
            ],
            charts=[],
            sections=[],
        )


def test_P0_DOCS_OUT_formula_over_text_cells_is_refused():
    with pytest.raises(ValidationError):
        spec(
            tables=[
                {
                    "id": "t1",
                    "columns": ["Region", "Units"],
                    "rows": [["N", "many"]],
                    "formulas": [{"column": "X", "expr": "[Units] * 2"}],
                }
            ],
            charts=[],
            sections=[],
        )


# ---- Excel ---------------------------------------------------------------------------------------------------------


def test_P0_DOCS_OUT_xlsx_has_real_formulas_with_cached_values_and_styled_header():
    data = render_xlsx(spec())
    sheet = next(s for s in load_workbook(io.BytesIO(data)).worksheets if s["A1"].value == "Region")
    assert sheet["D1"].value == "Revenue"
    assert sheet["D2"].value == "=B2*C2" and sheet["D3"].value == "=B3*C3"
    assert sheet["D5"].value == "=SUM(D2:D4)"
    assert sheet["A1"].font.bold and sheet["A1"].fill.fgColor.rgb not in (None, "00000000")
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:D4"
    cached = next(s for s in load_workbook(io.BytesIO(data), data_only=True).worksheets if s["A1"].value == "Region")
    assert [cached[f"D{r}"].value for r in (2, 3, 4, 5)] == [31.5, 81.0, 40, 152.5]
    assert cached["B5"].value == 12


def test_P0_DOCS_OUT_xlsx_has_a_native_chart_for_each_chart_block():
    with zipfile.ZipFile(io.BytesIO(render_xlsx(spec()))) as package:
        charts = [n for n in package.namelist() if n.startswith("xl/charts/chart")]
        assert len(charts) == 1
        assert b"Revenue" in package.read(charts[0])


def test_P0_DOCS_OUT_xlsx_hostile_text_is_still_a_string_and_the_file_is_deterministic():
    s = spec(
        tables=[{"id": "t1", "columns": ["Name"], "rows": [["=cmd|' /C calc'!A0"], ["@SUM(1)"]]}],
        charts=[],
        sections=[],
    )
    data = render_xlsx(s)
    sheet = load_workbook(io.BytesIO(data)).worksheets[-1]
    assert [c.data_type for c in sheet["A"][1:3]] == ["s", "s"]
    assert render_xlsx(s) == data


# ---- PDF -----------------------------------------------------------------------------------------------------------


def test_P0_DOCS_OUT_pdf_has_a_cover_page_then_the_report_with_page_numbers():
    data = render_pdf(spec())
    reader = PdfReader(io.BytesIO(data))
    assert len(reader.pages) >= 2
    cover = reader.pages[0].extract_text()
    assert "Q3 Sales Review" in cover and "North and South regions" in cover and "2026-10-20" in cover
    assert "Summary" not in cover  # the first section starts on the page after the cover
    body = reader.pages[1].extract_text()
    assert "Summary" in body and "North grew" in body and "Page 2" in body
    assert reader.metadata is not None and reader.metadata.title == "Q3 Sales Review"
    assert render_pdf(spec()) == data


def test_P0_DOCS_OUT_one_page_layout_has_no_cover_and_fits_a_flyer_on_one_page():
    s = spec(
        layout="one_page",
        tables=[],
        charts=[],
        sections=[
            {"heading": "Come along", "blocks": [{"type": "bullets", "bullets": ["Saturday 10am", "Free entry"]}]}
        ],
    )
    reader = PdfReader(io.BytesIO(render_pdf(s)))
    assert len(reader.pages) == 1
    text = reader.pages[0].extract_text()
    assert "Q3 Sales Review" in text and "Come along" in text and "Free entry" in text


def test_P0_DOCS_OUT_unknown_layout_is_refused():
    with pytest.raises(ValidationError):
        spec(layout="poster-with-script")


def test_P0_DOCS_OUT_pdf_table_numbers_use_thousands_separators():
    s = spec(
        tables=[{"id": "t1", "columns": ["Item", "Cost"], "rows": [["Rig", 1234567]]}],
        charts=[],
        sections=[{"heading": "Costs", "blocks": [{"type": "table", "table_id": "t1"}]}],
    )
    assert "1,234,567" in PdfReader(io.BytesIO(render_pdf(s))).pages[1].extract_text()


# ---- Word ----------------------------------------------------------------------------------------------------------


def _parts(data: bytes) -> dict[str, ET.Element]:
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        assert package.testzip() is None
        return {n: ET.fromstring(package.read(n)) for n in package.namelist() if n.endswith((".xml", ".rels"))}


def _style(p: ET.Element) -> str | None:
    node = p.find(f"./{W}pPr/{W}pStyle")
    return None if node is None else node.get(f"{W}val")


def test_P0_DOCS_OUT_docx_uses_real_word_styles_for_title_headings_and_bullets():
    parts = _parts(render_docx(spec()))
    styles = {s.get(f"{W}styleId") for s in parts["word/styles.xml"].iter(f"{W}style")}
    assert {"Title", "Subtitle", "Heading1", "ListBullet", "TableGrid"} <= styles
    paragraphs = list(parts["word/document.xml"].iter(f"{W}p"))
    by_style: dict[str | None, list[str]] = {}
    for p in paragraphs:
        by_style.setdefault(_style(p), []).append("".join(t.text or "" for t in p.iter(f"{W}t")))
    assert by_style["Title"] == ["Q3 Sales Review"]
    assert by_style["Heading1"] == ["Summary", "Next steps"]
    assert by_style["ListBullet"] == ["North grew", "South led"]


def test_P0_DOCS_OUT_docx_table_has_a_repeating_shaded_header_and_a_page_number_footer():
    parts = _parts(render_docx(spec()))
    table = parts["word/document.xml"].find(f".//{W}tbl")
    assert table is not None
    first = table.find(f"{W}tr")
    assert first is not None and first.find(f"{W}trPr/{W}tblHeader") is not None
    assert first.find(f"{W}tc/{W}tcPr/{W}shd") is not None
    footers = [n for n in parts if n.startswith("word/footer")]
    assert footers and any("PAGE" in f.get(f"{W}instr", "") for n in footers for f in parts[n].iter(f"{W}fldSimple"))
    assert parts["docProps/core.xml"].findtext("{http://purl.org/dc/elements/1.1/}title") == "Q3 Sales Review"


def test_P0_DOCS_OUT_docx_parses_with_python_docx_when_installed():
    docx = pytest.importorskip("docx")
    document = docx.Document(io.BytesIO(render_docx(spec())))
    assert [p.text for p in document.paragraphs if p.style.name == "Heading 1"] == ["Summary", "Next steps"]
    assert document.tables[0].rows[0].cells[0].text == "Region"
    assert document.tables[0].rows[1].cells[3].text == "31.5"
