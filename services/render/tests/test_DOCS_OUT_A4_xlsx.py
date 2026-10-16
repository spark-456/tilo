"""DOCS-OUT A4: an Excel deliverable made from the tables an agent produced."""

import io
from datetime import date

import pytest
from openpyxl import load_workbook
from tilo_render.spec import DeliverableSpec
from tilo_render.xlsx import MAX_XLSX_CELLS, render_xlsx


def spec(**extra) -> DeliverableSpec:
    body = {
        "title": "Sales",
        "date": "2026-10-16",
        "tables": [
            {
                "id": "t1",
                "columns": ["Region", "Closed on", "Units", "Revenue"],
                "rows": [
                    ["North", "2026-09-30", 3, 10.5],
                    ["South", "2026-10-02", 4, 20.25],
                ],
            }
        ],
        "sections": [
            {
                "heading": "Summary",
                "blocks": [{"type": "paragraph", "text": "Revenue is up."}, {"type": "table", "table_id": "t1"}],
            }
        ],
    }
    body.update(extra)
    return DeliverableSpec.model_validate(body)


def book(s: DeliverableSpec):
    return load_workbook(io.BytesIO(render_xlsx(s)))


def test_DOCS_OUT_A4_table_becomes_a_sheet_with_header_types_widths_and_frozen_header():
    wb = book(spec())
    assert wb.sheetnames[0] == "Summary"  # the paragraph sits on its own sheet ahead of the table
    table = next(s for s in wb.worksheets if s["A1"].value == "Region")
    assert [c.value for c in table[1]] == ["Region", "Closed on", "Units", "Revenue"]
    assert table["A1"].font.bold
    assert table.freeze_panes == "A2"
    assert table["C2"].value == 3 and table["C2"].data_type == "n"
    assert table["D3"].value == 20.25
    assert table["B2"].value.date() == date(2026, 9, 30) and "yyyy" in table["B2"].number_format
    assert table.column_dimensions["A"].width >= 10
    assert table.auto_filter.ref == "A1:D3"


def test_DOCS_OUT_A4_numeric_columns_get_a_totals_row_of_internal_sums_only():
    table = next(s for s in book(spec()).worksheets if s["A1"].value == "Region")
    assert table["A4"].value == "Total" and table["A4"].font.bold
    assert table["C4"].value == "=SUM(C2:C3)"
    assert table["D4"].value == "=SUM(D2:D3)"
    assert table["B4"].value is None  # dates are not summed


def test_DOCS_OUT_A4_model_text_is_data_never_a_formula():
    hostile = '=HYPERLINK("http://x","c")'
    cmd = "=cmd|' /C calc'!A0"
    s = spec(
        tables=[
            {
                "id": "t1",
                "columns": ["Name", "Note"],
                "rows": [[hostile, "+SUM(1)"], [cmd, "@A1"]],
            }
        ],
        sections=[{"heading": "H", "blocks": [{"type": "table", "table_id": "t1"}]}],
    )
    sheet = next(s for s in book(s).worksheets if s["A1"].value == "Name")
    for ref, text in (("A2", hostile), ("B2", "+SUM(1)"), ("A3", cmd), ("B3", "@A1")):
        assert sheet[ref].data_type == "s" and sheet[ref].value == text
    assert all(c.data_type != "f" for row in sheet.iter_rows() for c in row)


def test_DOCS_OUT_A4_no_external_links_in_the_package():
    import zipfile

    with zipfile.ZipFile(io.BytesIO(render_xlsx(spec()))) as package:
        names = package.namelist()
        assert not any("externalLink" in n or "vba" in n.lower() or n.endswith(".bin") for n in names)


def test_DOCS_OUT_A4_report_without_tables_still_makes_a_readable_sheet():
    s = DeliverableSpec.model_validate(
        {
            "title": "Notes",
            "date": "2026-10-16",
            "sections": [{"heading": "One", "blocks": [{"type": "bullets", "bullets": ["a", "b"]}]}],
        }
    )
    sheet = book(s).worksheets[0]
    values = [c.value for row in sheet.iter_rows() for c in row if c.value]
    assert values[0] == "Notes" and "One" in values and "- a" in values and "- b" in values


def test_DOCS_OUT_A4_sheet_names_are_valid_and_unique():
    tables = [{"id": f"t{i}", "columns": ["A"], "rows": [[i]]} for i in range(3)]
    blocks = [{"type": "table", "table_id": f"t{i}"} for i in range(3)]
    s = spec(
        tables=tables,
        sections=[{"heading": "Q?/[1]: " + "x" * 40, "blocks": blocks}],
    )
    names = book(s).sheetnames
    assert len(set(n.lower() for n in names)) == len(names)
    assert all(len(n) <= 31 and not any(c in n for c in "[]:*?/\\") for n in names)


def test_DOCS_OUT_A4_cell_cap_refuses_instead_of_producing_a_huge_workbook():
    rows = [[f"r{i}", i] for i in range(500)]
    tables = [{"id": f"t{i}", "columns": ["a", "b"], "rows": rows} for i in range(8)]
    assert MAX_XLSX_CELLS < 8000
    with pytest.raises(ValueError):
        render_xlsx(spec(tables=tables, sections=[]))


def test_DOCS_OUT_A4_output_is_a_zip_workbook():
    data = render_xlsx(spec())
    assert data.startswith(b"PK\x03\x04")
