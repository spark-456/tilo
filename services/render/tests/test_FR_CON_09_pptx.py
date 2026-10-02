"""Acceptance tests for AGR-02 PPTX renderer with native charts (FR-CON-09)."""

import io
import json
from pathlib import Path

import pptx
from pptx.dml.color import RGBColor
from tilo_render.pptx import render_pptx
from tilo_render.spec import (
    BulletsBlock,
    Chart,
    ChartBlock,
    ChartKind,
    DeliverableSpec,
    ParagraphBlock,
    Section,
    Table,
    TableBlock,
)


def _make_sample_spec(rows: list[list[str | int | float]] | None = None) -> DeliverableSpec:
    default_rows: list[list[str | int | float]] = [["Alpha", 10.0, 5.0], ["Beta", 20.0, 15.0], ["Gamma", 30.0, 25.0]]
    t1 = Table(id="t1", columns=["label", "s1", "s2"], rows=rows or default_rows)
    return DeliverableSpec(
        title="Sample Report",
        subtitle="Subtitle text",
        date="2026-10-02",
        tables=[t1],
        charts=[
            Chart(id="c_bar", kind=ChartKind.BAR, table_id="t1", x="label", y=["s1", "s2"], title="Bar Chart"),
            Chart(id="c_line", kind=ChartKind.LINE, table_id="t1", x="label", y="s1", title="Line Chart"),
            Chart(id="c_pie", kind=ChartKind.PIE, table_id="t1", x="label", y="s2", title="Pie Chart"),
        ],
        sections=[
            Section(
                heading="Overview",
                blocks=[
                    ParagraphBlock(text="Intro paragraph."),
                    BulletsBlock(bullets=["Point 1", "Point 2"]),
                    TableBlock(table_id="t1"),
                    ChartBlock(chart_id="c_bar"),
                    ChartBlock(chart_id="c_line"),
                    ChartBlock(chart_id="c_pie"),
                ],
            )
        ],
    )


def test_FR_CON_09_pptx_chart_values_equal_the_source_table() -> None:
    spec = _make_sample_spec()
    prs = pptx.Presentation(io.BytesIO(render_pptx(spec)))

    chart_slides = [s for s in prs.slides if any(sh.has_chart for sh in s.shapes)]
    assert len(chart_slides) == 3

    bar_chart = next(sh.chart for sh in chart_slides[0].shapes if sh.has_chart)
    series_map = {series.name: list(series.values) for series in bar_chart.series}
    assert series_map["s1"] == [10.0, 20.0, 30.0] and series_map["s2"] == [5.0, 15.0, 25.0]

    line_chart = next(sh.chart for sh in chart_slides[1].shapes if sh.has_chart)
    assert [s.values for s in line_chart.series][0] == (10.0, 20.0, 30.0)

    pie_chart = next(sh.chart for sh in chart_slides[2].shapes if sh.has_chart)
    assert [s.values for s in pie_chart.series][0] == (5.0, 15.0, 25.0)


def test_FR_CON_09_pptx_opens_and_has_one_slide_per_block() -> None:
    prs = pptx.Presentation(io.BytesIO(render_pptx(_make_sample_spec())))
    assert len(prs.slides) == 8

    long_rows: list[list[str | int | float]] = [[f"Item {i}", float(i), float(i * 2)] for i in range(25)]
    prs_long = pptx.Presentation(io.BytesIO(render_pptx(_make_sample_spec(rows=long_rows))))
    assert len(prs_long.slides) == 10

    table_slides = [s for s in prs_long.slides if any(sh.has_table for sh in s.shapes)]
    assert len(table_slides) == 3
    for s in table_slides:
        tbl = next(sh.table for sh in s.shapes if sh.has_table)
        assert [tbl.cell(0, col).text for col in range(3)] == ["label", "s1", "s2"]


def test_FR_CON_09_pptx_colours_come_from_token_roles() -> None:
    prs = pptx.Presentation(io.BytesIO(render_pptx(_make_sample_spec())))
    tokens = json.loads(Path("design/tokens.json").read_text(encoding="utf-8"))
    brand_primary = tokens["color"]["brand"]["primary"]["value"].lstrip("#")
    assert prs.slides[0].background.fill.fore_color.rgb == RGBColor.from_string(brand_primary)


def test_FR_CON_09_pptx_escapes_untrusted_text() -> None:
    evil_text = ("<script>alert('xss')</script> & 'quotes' \x00\x07\x1b " * 800)[:20000]
    expected_clean = "".join(
        ch for ch in evil_text if ch in "\t\n\r" or (ord(ch) >= 32 and ord(ch) != 127 and not (0x80 <= ord(ch) <= 0x9F))
    )
    spec = DeliverableSpec(
        title="Escaping <Test> & More \x00\x07",
        date="2026-10-02",
        sections=[
            Section(
                heading="Sec <1> & Info \x07",
                blocks=[ParagraphBlock(text=evil_text), BulletsBlock(bullets=["Point & <tag> \x00"])],
            )
        ],
    )
    prs = pptx.Presentation(io.BytesIO(render_pptx(spec)))
    text_content = " ".join(sh.text for sh in prs.slides[2].shapes if sh.has_text_frame)
    assert expected_clean in text_content
    assert "\x00" not in text_content and "\x07" not in text_content
    assert "<script>" in text_content and "&" in text_content
