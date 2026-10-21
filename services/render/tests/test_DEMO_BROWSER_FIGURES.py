"""DEMO-BROWSER (owner, 21 Oct: "take screenshots and save it as pdf with proper captioning"): a deliverable may place
pictures the run recorded (Tilo's browser frames) as figures. The writer writes only {"type": "figure", "figure_id",
"caption"}; code supplies the spec's `figures` (recorded PNG or JPEG bytes and the page's address). PDF and Word
print the picture at the text width with "Figure n: caption" and "Source: <address>" under it; slides get an image
slide. Pictures stay within fixed caps (count, bytes, pixels) and must be real images.
"""

from __future__ import annotations

import base64
import io
import zipfile

import pytest
from PIL import Image
from pydantic import ValidationError
from tilo_render.spec import MAX_FIGURES, DeliverableSpec


def _png(width: int = 600, height: int = 400, colour: str = "#2a6fb0") -> str:
    out = io.BytesIO()
    Image.new("RGB", (width, height), colour).save(out, "PNG")
    return base64.b64encode(out.getvalue()).decode("ascii")


SOURCE = "https://en.wikipedia.org/wiki/List_of_tallest_buildings_in_India"


def _spec(**over):
    data = {
        "title": "Tallest buildings",
        "date": "2026-10-21",
        "tables": [{"id": "t1", "columns": ["Country", "Building", "Height (m)"], "rows": [["India", "Palais", 320]]}],
        "figures": [{"id": "F1", "image_base64": _png(), "source": SOURCE}],
        "sections": [
            {
                "heading": "Results",
                "blocks": [
                    {"type": "table", "table_id": "t1"},
                    {
                        "type": "figure",
                        "figure_id": "F1",
                        "caption": "Wikipedia's list of the tallest buildings in India",
                    },
                ],
            }
        ],
    }
    data.update(over)
    return DeliverableSpec.model_validate(data)


def test_DEMO_BROWSER_a_figure_prints_in_the_pdf_with_its_caption_and_source():
    from pypdf import PdfReader
    from tilo_render.pdf import render_pdf

    reader = PdfReader(io.BytesIO(render_pdf(_spec())))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Figure 1: Wikipedia's list of the tallest buildings in India" in text
    assert "Source: https://en.wikipedia.org/wiki/List_of_tallest_buildings_in_India" in text
    images = [image for page in reader.pages for image in page.images]
    assert len(images) == 1 and images[0].image.size == (600, 400)


def test_DEMO_BROWSER_a_figure_is_a_picture_in_the_word_document():
    from tilo_render.docx import render_docx

    with zipfile.ZipFile(io.BytesIO(render_docx(_spec()))) as package:
        names = package.namelist()
        document = package.read("word/document.xml").decode()
        rels = package.read("word/_rels/document.xml.rels").decode()
        types = package.read("[Content_Types].xml").decode()
        picture = package.read("word/media/figure1.png")
    assert "word/media/figure1.png" in names and Image.open(io.BytesIO(picture)).size == (600, 400)
    assert 'r:embed="rIdFigure1"' in document and "media/figure1.png" in rels and 'Extension="png"' in types
    assert "Figure 1: Wikipedia" in document and f"Source: {SOURCE}" in document


def test_DEMO_BROWSER_a_figure_is_an_image_slide():
    from pptx import Presentation
    from tilo_render.pptx import render_pptx

    deck = Presentation(io.BytesIO(render_pptx(_spec())))
    pictures = [shape for slide in deck.slides for shape in slide.shapes if shape.shape_type == 13]
    assert len(pictures) == 1


@pytest.mark.parametrize(
    "over",
    [
        {"figures": []},  # a block naming a picture the spec does not carry
        {"figures": [{"id": "F1", "image_base64": base64.b64encode(b"not a picture").decode(), "source": SOURCE}]},
        {"figures": [{"id": "F1", "image_base64": _png(3000, 2000), "source": SOURCE}]},  # over the pixel cap
        {"figures": [{"id": "F1", "image_base64": _png(), "source": "line one\nline two"}]},
    ],
    ids=["unknown id", "not an image", "too many pixels", "two-line source"],
)
def test_DEMO_BROWSER_a_figure_that_is_not_a_real_recorded_picture_is_refused(over):
    with pytest.raises(ValidationError):
        _spec(**over)


def test_DEMO_BROWSER_figures_are_capped():
    many = [{"id": f"F{n}", "image_base64": _png(40, 30), "source": SOURCE} for n in range(MAX_FIGURES + 1)]
    with pytest.raises(ValidationError):
        _spec(figures=many)
