"""DOCS-OUT A4: fill a user's uploaded Excel template with values the writer returned, keeping its formatting."""

import io
import zipfile
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from tilo_render.xlsx import MAX_FILL_CELLS, FillCell, describe_workbook, fill_workbook


def template() -> bytes:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Invoice"
    sheet["A1"], sheet["B1"], sheet["C1"], sheet["D1"] = "Item", "Qty", "Due", "Price"
    for ref in ("A1", "B1", "C1", "D1"):
        sheet[ref].font = Font(bold=True, color="FF0000")
        sheet[ref].fill = PatternFill("solid", start_color="FFFF00", end_color="FFFF00")
    sheet["D2"].number_format = "$#,##0.00"
    sheet["C2"].number_format = "dd/mm/yyyy"
    sheet["A5"], sheet["D5"] = "Total", "=SUM(D2:D4)"
    sheet.merge_cells("A7:D7")
    sheet["A7"] = "Notes"
    sheet.column_dimensions["A"].width = 33
    sheet.freeze_panes = "A2"
    wb.create_sheet("Other")["A1"] = "keep"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def cells(*items: tuple[str, str, object]) -> list[FillCell]:
    return [FillCell(sheet=s, cell=c, value=v) for s, c, v in items]  # type: ignore[arg-type]


def load(data: bytes):
    return load_workbook(io.BytesIO(data))


def test_DOCS_OUT_A4_fill_writes_values_and_keeps_the_templates_formatting():
    out = fill_workbook(
        template(),
        cells(
            ("Invoice", "A2", "Widget"), ("Invoice", "B2", 3), ("Invoice", "D2", 9.5), ("Invoice", "C2", "2026-11-01")
        ),
    )
    sheet = load(out)["Invoice"]
    assert (sheet["A2"].value, sheet["B2"].value, sheet["D2"].value) == ("Widget", 3, 9.5)
    assert sheet["C2"].value.date() == date(2026, 11, 1) and sheet["C2"].number_format == "dd/mm/yyyy"
    assert sheet["D2"].number_format == "$#,##0.00"
    assert sheet["A1"].font.bold and sheet["A1"].fill.start_color.rgb.endswith("FFFF00")
    assert sheet.column_dimensions["A"].width == 33 and sheet.freeze_panes == "A2"
    assert sheet["D5"].value == "=SUM(D2:D4)" and "A7:D7" in {str(r) for r in sheet.merged_cells.ranges}
    assert load(out)["Other"]["A1"].value == "keep"


def test_DOCS_OUT_A4_filled_text_is_data_never_a_formula():
    out = fill_workbook(template(), cells(("Invoice", "A2", '=HYPERLINK("http://x","c")'), ("Invoice", "A3", "+1")))
    sheet = load(out)["Invoice"]
    assert sheet["A2"].data_type == "s" and sheet["A2"].value == '=HYPERLINK("http://x","c")'
    assert sheet["A3"].data_type == "s"


@pytest.mark.parametrize(
    "item",
    [
        ("Missing", "A2", 1),  # no such sheet: never created
        ("Invoice", "D5", 1),  # a template formula is never overwritten
        ("Invoice", "B7", 1),  # inside a merged range, not its anchor
        ("Invoice", "ZZZZ1", 1),
        ("Invoice", "A0", 1),
        ("Invoice", "A100001", 1),
        ("Invoice", "$A$1", 1),
        ("Invoice", "A1:B2", 1),
    ],
)
def test_DOCS_OUT_A4_unsafe_targets_are_refused(item):
    with pytest.raises(ValueError):
        fill_workbook(template(), cells(item))


def test_DOCS_OUT_A4_fill_is_capped_and_values_are_bounded():
    many = [FillCell(sheet="Invoice", cell=f"A{i + 2}", value=i) for i in range(MAX_FILL_CELLS + 1)]
    with pytest.raises(ValueError):
        fill_workbook(template(), many)
    with pytest.raises(ValueError):
        fill_workbook(template(), cells(("Invoice", "A2", "x" * 5000)))
    with pytest.raises(ValueError):
        fill_workbook(template(), cells(("Invoice", "A2", float("inf"))))


@pytest.mark.parametrize("part", ["xl/externalLinks/externalLink1.xml", "xl/vbaProject.bin", "xl/embeddings/o.bin"])
def test_DOCS_OUT_A4_templates_with_external_or_active_parts_are_refused(part):
    source = io.BytesIO(template())
    patched = io.BytesIO()
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(patched, "w") as zout:
        for info in zin.infolist():
            zout.writestr(info, zin.read(info))
        zout.writestr(part, b"x")
    with pytest.raises(ValueError):
        fill_workbook(patched.getvalue(), cells(("Invoice", "A2", 1)))
    with pytest.raises(ValueError):
        describe_workbook(patched.getvalue())


def test_DOCS_OUT_A4_not_a_workbook_is_refused():
    with pytest.raises(ValueError):
        fill_workbook(b"PK\x03\x04not a zip", cells(("Invoice", "A2", 1)))


def test_DOCS_OUT_A4_describe_lists_sheets_labels_and_addresses_within_a_cap():
    text = describe_workbook(template())
    assert "Invoice" in text and "A1=Item" in text and "D5" in text and "Other" in text
    big = Workbook()
    for i in range(1, 400):
        big.active.cell(row=i, column=1, value=f"label {i}")
    buffer = io.BytesIO()
    big.save(buffer)
    assert len(describe_workbook(buffer.getvalue(), max_chars=1500)) <= 1500
